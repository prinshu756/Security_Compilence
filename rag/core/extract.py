"""Extract plain text from every source type in knowledge_base."""

from pathlib import Path


def extract_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n\n".join(parts)


def extract_docx(path: Path) -> str:
    import io
    import docx
    doc = docx.Document(io.BytesIO(path.read_bytes()))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(", ".join(cells))
    return "\n".join(parts)


def extract_http(url: str) -> str:
    import re
    import requests
    from bs4 import BeautifulSoup
    resp = requests.get(url, timeout=20)
    resp.raise_for_status()
    soup = BeautifulSoup(resp.text, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    text = soup.get_text(separator="\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    return path.read_text(encoding="utf-8", errors="ignore")